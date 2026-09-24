#!/usr/bin/env python3

from __future__ import annotations

import argparse
import csv
import gzip
import io
import logging
import os
import tarfile
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from functools import partial
from pathlib import Path
from threading import Lock, local
from typing import Callable

import boto3
import pandas as pd
from google.cloud import storage
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

try:
    from tqdm.auto import tqdm as tqdm_progress
except ImportError:  # pragma: no cover - exercised only when tqdm is unavailable.
    tqdm_progress = None


DEFAULT_ARTIFACT_DIR = Path("testdata")
OVERSIZED_TARGET_BYTES = 52 * 1024 * 1024
LARGEFILE_TARGET_BYTES = 47 * 1024 * 1024

LOREM_TEXT = """
Lorem ipsum dolor sit amet, consectetur adipiscing elit.
Sed do eiusmod tempor incididunt ut labore et dolore magna aliqua.
Ut enim ad minim veniam, quis nostrud exercitation ullamco laboris.
Duis aute irure dolor in reprehenderit in voluptate velit esse cillum dolore.
Excepteur sint occaecat cupidatat non proident, sunt in culpa qui officia deserunt.
""".strip()

PII_ROWS = [
    {
        "name": "Ananya Rao",
        "email": "ananya.rao@example.com",
        "phone": "+1-415-555-0134",
        "ssn": "123-45-6789",
        "note": "Customer requested invoice copy.",
    },
    {
        "name": "Rahul Mehta",
        "email": "rahul.mehta@example.com",
        "phone": "+91-98765-43210",
        "ssn": "987-65-4321",
        "note": "Account verification pending.",
    },
]

NO_PII_ROWS = [
    {"col1": "lorem", "col2": "ipsum", "col3": "dolor"},
    {"col1": "sit", "col2": "amet", "col3": "consectetur"},
]

LOGGER = logging.getLogger("fixture_generator")
UPLOAD_CLIENTS = local()
FALLBACK_PROGRESS_LOCK = Lock()


GeneratorFunc = Callable[[Path, str, int | None], None]


@dataclass(frozen=True)
class FixtureTask:
    path: Path
    action: Callable[[], None]


@dataclass(frozen=True)
class UploadTarget:
    provider: str
    bucket: str

    @property
    def display_name(self) -> str:
        scheme = "s3" if self.provider == "s3" else "gs"
        return f"{scheme}://{self.bucket}"


@dataclass(frozen=True)
class GenerationResult:
    generated: list[Path]
    skipped: list[Path]


class SimpleProgressBar:
    def __init__(self, total: int, desc: str, unit: str, position: int = 0) -> None:
        self.total = total
        self.desc = desc
        self.unit = unit
        self.position = position
        self.current = 0
        self.closed = False
        progress_write(f"{self.desc}: 0/{self.total} {self.unit}")

    def __enter__(self) -> SimpleProgressBar:
        return self

    def __exit__(self, exc_type, exc, tb) -> None:
        self.close()

    def update(self, count: int = 1) -> None:
        if self.closed:
            return
        self.current += count
        progress_write(f"{self.desc}: {self.current}/{self.total} {self.unit}")

    def close(self) -> None:
        if self.closed:
            return
        self.closed = True
        progress_write(f"{self.desc}: complete ({self.current}/{self.total} {self.unit})")


def progress_write(message: str) -> None:
    if tqdm_progress is not None:
        tqdm_progress.write(message)
        return

    with FALLBACK_PROGRESS_LOCK:
        print(message, flush=True)


class TqdmLoggingHandler(logging.Handler):
    def emit(self, record: logging.LogRecord) -> None:
        try:
            progress_write(self.format(record))
        except Exception:
            self.handleError(record)


def setup_logging() -> None:
    root_logger = logging.getLogger()
    if root_logger.handlers:
        root_logger.handlers.clear()

    handler = TqdmLoggingHandler()
    handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s", "%Y-%m-%d %H:%M:%S"))

    root_logger.addHandler(handler)
    root_logger.setLevel(logging.INFO)

    if tqdm_progress is None:
        LOGGER.warning("tqdm is not installed; using simplified progress output instead.")


def make_progress(total: int, desc: str, unit: str, position: int = 0):
    if tqdm_progress is None:
        return SimpleProgressBar(total=total, desc=desc, unit=unit, position=position)

    return tqdm_progress(total=total, desc=desc, unit=unit, position=position, dynamic_ncols=True, leave=True)


def ensure_dir(base_dir: Path, folder: str) -> Path:
    path = base_dir / folder
    path.mkdir(parents=True, exist_ok=True)
    return path


def pii_text() -> str:
    return "\n".join(
        f"Name: {r['name']}\nEmail: {r['email']}\nPhone: {r['phone']}\nSSN: {r['ssn']}\nNote: {r['note']}\n"
        for r in PII_ROWS
    )


def no_pii_text() -> str:
    return LOREM_TEXT


def rows_for(kind: str) -> list[dict[str, str]]:
    return PII_ROWS if kind == "pii" else NO_PII_ROWS


def inflate_plain_file(path: Path, target_bytes: int) -> None:
    chunk = (LOREM_TEXT + "\n").encode("utf-8") * 1024

    with path.open("ab") as file_handle:
        while path.stat().st_size < target_bytes:
            file_handle.write(chunk)


def inflate_binary_file(path: Path, target_bytes: int) -> None:
    chunk = os.urandom(1024 * 1024)

    with path.open("ab") as file_handle:
        while path.stat().st_size < target_bytes:
            remaining = target_bytes - path.stat().st_size
            file_handle.write(chunk[: min(len(chunk), remaining)])


def make_zip_based_file_large(path: Path, target_bytes: int) -> None:
    payload_size = max(target_bytes - path.stat().st_size, 1)
    chunk = os.urandom(1024 * 1024)

    with zipfile.ZipFile(path, "a", compression=zipfile.ZIP_STORED) as zip_handle:
        with zip_handle.open("customXml/large_payload.bin", "w") as file_handle:
            written = 0
            while written < payload_size:
                remaining = payload_size - written
                data = chunk[: min(len(chunk), remaining)]
                file_handle.write(data)
                written += len(data)


def generate_txt(path: Path, kind: str, target_bytes: int | None = None) -> None:
    path.write_text(pii_text() if kind == "pii" else no_pii_text(), encoding="utf-8")
    if target_bytes:
        inflate_plain_file(path, target_bytes)


def generate_log(path: Path, kind: str, target_bytes: int | None = None) -> None:
    if kind == "pii":
        lines = [
            "INFO user='Ananya Rao' email='ananya.rao@example.com' phone='+1-415-555-0134' ssn='123-45-6789'",
            "WARN user='Rahul Mehta' email='rahul.mehta@example.com' phone='+91-98765-43210' ssn='987-65-4321'",
        ]
    else:
        lines = [
            "INFO Lorem ipsum dolor sit amet",
            "DEBUG Consectetur adipiscing elit",
            "WARN Sed do eiusmod tempor incididunt",
        ]

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    if target_bytes:
        inflate_plain_file(path, target_bytes)


def generate_csv(path: Path, kind: str, target_bytes: int | None = None) -> None:
    rows = rows_for(kind)

    with path.open("w", newline="", encoding="utf-8") as file_handle:
        writer = csv.DictWriter(file_handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)

        if target_bytes:
            while path.stat().st_size < target_bytes:
                writer.writerow(rows[0])


def generate_pdf(path: Path, kind: str, target_bytes: int | None = None) -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter)
    text = pdf.beginText(72, 720)
    text.setFont("Helvetica", 11)

    for line in (pii_text() if kind == "pii" else no_pii_text()).splitlines():
        text.textLine(line)

    pdf.drawText(text)
    pdf.showPage()
    pdf.save()

    if target_bytes:
        inflate_plain_file(path, target_bytes)


def generate_docx(path: Path, kind: str, target_bytes: int | None = None) -> None:
    document = Document()
    document.add_heading("Scanner Fixture Document", level=1)

    for line in (pii_text() if kind == "pii" else no_pii_text()).splitlines():
        document.add_paragraph(line)

    document.save(path)

    if target_bytes:
        make_zip_based_file_large(path, target_bytes)


def generate_pptx(path: Path, kind: str, target_bytes: int | None = None) -> None:
    presentation = Presentation()

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Scanner Fixture Presentation"
    slide.placeholders[1].text = pii_text() if kind == "pii" else no_pii_text()

    presentation.save(path)

    if target_bytes:
        make_zip_based_file_large(path, target_bytes)


def generate_xlsx(path: Path, kind: str, target_bytes: int | None = None) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "fixture"

    rows = rows_for(kind)
    headers = list(rows[0].keys())

    sheet.append(headers)
    for row in rows:
        sheet.append([row[header] for header in headers])

    workbook.save(path)

    if target_bytes:
        make_zip_based_file_large(path, target_bytes)


def generate_parquet(path: Path, kind: str, target_bytes: int | None = None) -> None:
    rows = rows_for(kind)
    dataframe = pd.DataFrame(rows * 100)
    dataframe.to_parquet(path, engine="pyarrow", index=False)

    if target_bytes:
        repeat = 100_000
        previous_size = path.stat().st_size
        max_attempts = 8

        for _ in range(max_attempts):
            if previous_size >= target_bytes:
                break
            dataframe = pd.DataFrame(rows * repeat)
            dataframe.to_parquet(path, engine="pyarrow", index=False)
            current_size = path.stat().st_size
            if current_size <= previous_size:
                LOGGER.warning(
                    "Stopped growing parquet fixture %s at %s bytes because the file size no longer increased.",
                    path,
                    current_size,
                )
                break
            previous_size = current_size
            repeat *= 2


def generate_image(path: Path, image_format: str, kind: str, target_bytes: int | None = None) -> None:
    image = Image.new("RGB", (1200, 800), "white")
    draw = ImageDraw.Draw(image)

    text = pii_text() if kind == "pii" else no_pii_text()
    draw.text((40, 40), text, fill="black")

    save_kwargs = {}
    if image_format.upper() in {"JPEG", "WEBP"}:
        save_kwargs["quality"] = 95

    image.save(path, format=image_format.upper(), **save_kwargs)

    if target_bytes:
        inflate_binary_file(path, target_bytes)


def generate_binary(path: Path, kind: str, target_bytes: int | None = None) -> None:
    payload = os.urandom(4096)

    if kind == "pii":
        payload += (
            b"\nName: Ananya Rao\n"
            b"Email: ananya.rao@example.com\n"
            b"Phone: +1-415-555-0134\n"
            b"SSN: 123-45-6789\n"
        )
    else:
        payload += b"\nLorem ipsum dolor sit amet, consectetur adipiscing elit.\n"

    path.write_bytes(payload)

    if target_bytes:
        inflate_binary_file(path, target_bytes)


def generate_zip(path: Path, kind: str, target_bytes: int | None = None) -> None:
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zip_handle:
        zip_handle.writestr("sample.txt", pii_text() if kind == "pii" else no_pii_text())

        rows = rows_for(kind)
        csv_buffer = io.StringIO()
        writer = csv.DictWriter(csv_buffer, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
        zip_handle.writestr("sample.csv", csv_buffer.getvalue())

        zip_handle.writestr("sample.log", "INFO Lorem ipsum dolor sit amet\n")

        if target_bytes:
            zip_handle.writestr(
                "large_payload.txt",
                b"x" * target_bytes,
                compress_type=zipfile.ZIP_STORED,
            )


def generate_gz(path: Path, kind: str, target_bytes: int | None = None) -> None:
    content = pii_text() if kind == "pii" else no_pii_text()

    with gzip.open(path, "wt", encoding="utf-8") as file_handle:
        file_handle.write(content)

    if target_bytes:
        inflate_binary_file(path, target_bytes)


def generate_tar_gz(path: Path, kind: str, target_bytes: int | None = None) -> None:
    with tarfile.open(path, "w:gz") as tar_handle:
        files = {
            "sample.txt": pii_text() if kind == "pii" else no_pii_text(),
            "sample.log": "INFO Lorem ipsum dolor sit amet\n",
            "sample.csv": "col1,col2,col3\nlorem,ipsum,dolor\n",
        }

        for name, content in files.items():
            data = content.encode("utf-8")
            info = tarfile.TarInfo(name=name)
            info.size = len(data)
            tar_handle.addfile(info, io.BytesIO(data))

        if target_bytes:
            payload = os.urandom(target_bytes)
            info = tarfile.TarInfo(name="large_payload.bin")
            info.size = len(payload)
            tar_handle.addfile(info, io.BytesIO(payload))


def generate_nested_zip(path: Path, kind: str, target_bytes: int | None = None) -> None:
    inner = io.BytesIO()

    with zipfile.ZipFile(inner, "w", compression=zipfile.ZIP_DEFLATED) as inner_zip:
        inner_zip.writestr("inner_sample.txt", pii_text() if kind == "pii" else no_pii_text())

        if target_bytes:
            inner_zip.writestr(
                "inner_large_payload.bin",
                os.urandom(target_bytes),
                compress_type=zipfile.ZIP_STORED,
            )

    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as outer_zip:
        outer_zip.writestr("inner.zip", inner.getvalue(), compress_type=zipfile.ZIP_STORED)
        outer_zip.writestr("outer_manifest.txt", "Lorem ipsum dolor sit amet\n")


def generate_nested_gz(path: Path, kind: str, target_bytes: int | None = None) -> None:
    content = pii_text() if kind == "pii" else no_pii_text()
    first = gzip.compress(content.encode("utf-8"))
    second = gzip.compress(first)

    path.write_bytes(second)

    if target_bytes:
        inflate_binary_file(path, target_bytes)


def build_fixture_tasks(artifact_dir: Path) -> list[FixtureTask]:
    tasks: list[FixtureTask] = []

    specs: list[tuple[str, str, GeneratorFunc]] = [
        ("pdf", ".pdf", generate_pdf),
        ("docx", ".docx", generate_docx),
        ("pptx", ".pptx", generate_pptx),
        ("xlsx", ".xlsx", generate_xlsx),
        ("csv", ".csv", generate_csv),
        ("txt", ".txt", generate_txt),
        ("log", ".log", generate_log),
        ("parquet", ".parquet", generate_parquet),
        ("jpeg", ".jpg", lambda p, k, t: generate_image(p, "JPEG", k, t)),
        ("png", ".png", lambda p, k, t: generate_image(p, "PNG", k, t)),
        ("webp", ".webp", lambda p, k, t: generate_image(p, "WEBP", k, t)),
        ("zip", ".zip", generate_zip),
        ("gz", ".txt.gz", generate_gz),
        ("tar_gz", ".tar.gz", generate_tar_gz),
        ("nested_zip", ".zip", generate_nested_zip),
        ("nested_gz", ".csv.gz.gz", generate_nested_gz),
        ("bin", ".bin", generate_binary),
    ]

    variants = [
        ("has_pii_", "pii", None),
        ("has_no_pii_", "no_pii", None),
        ("largefile_", "no_pii", LARGEFILE_TARGET_BYTES),
        ("oversized_", "no_pii", OVERSIZED_TARGET_BYTES),
    ]

    for folder, extension, generator in specs:
        out_dir = artifact_dir / folder
        for prefix, kind, target_bytes in variants:
            path = out_dir / f"{prefix}sample{extension}"
            tasks.append(FixtureTask(path=path, action=partial(generator, path, kind, target_bytes)))

    no_ext_dir = artifact_dir / "no_extension"
    no_ext_specs: list[tuple[str, Callable[[Path], None]]] = [
        ("has_pii_jpeg_no_ext", lambda p: generate_image(p, "JPEG", "pii", None)),
        ("has_no_pii_jpeg_no_ext", lambda p: generate_image(p, "JPEG", "no_pii", None)),
        ("largefile_jpeg_no_ext", lambda p: generate_image(p, "JPEG", "no_pii", LARGEFILE_TARGET_BYTES)),
        ("oversized_jpeg_no_ext", lambda p: generate_image(p, "JPEG", "no_pii", OVERSIZED_TARGET_BYTES)),
        ("has_pii_csv_no_ext", lambda p: generate_csv(p, "pii", None)),
        ("has_no_pii_csv_no_ext", lambda p: generate_csv(p, "no_pii", None)),
        ("largefile_csv_no_ext", lambda p: generate_csv(p, "no_pii", LARGEFILE_TARGET_BYTES)),
        ("oversized_csv_no_ext", lambda p: generate_csv(p, "no_pii", OVERSIZED_TARGET_BYTES)),
        ("has_pii_unknown_binary_no_ext", lambda p: generate_binary(p, "pii", None)),
        ("has_no_pii_unknown_binary_no_ext", lambda p: generate_binary(p, "no_pii", None)),
        ("largefile_unknown_binary_no_ext", lambda p: generate_binary(p, "no_pii", LARGEFILE_TARGET_BYTES)),
        ("oversized_unknown_binary_no_ext", lambda p: generate_binary(p, "no_pii", OVERSIZED_TARGET_BYTES)),
    ]

    for name, generator in no_ext_specs:
        path = no_ext_dir / name
        tasks.append(FixtureTask(path=path, action=partial(generator, path)))

    return tasks


def generate_missing_fixtures(artifact_dir: Path) -> GenerationResult:
    tasks = build_fixture_tasks(artifact_dir)
    generated: list[Path] = []
    skipped: list[Path] = []

    artifact_dir.mkdir(parents=True, exist_ok=True)

    with make_progress(total=len(tasks), desc="Generate fixtures", unit="file") as progress:
        for task in tasks:
            if task.path.exists():
                LOGGER.info("Skipping existing file: %s", task.path)
                skipped.append(task.path)
            else:
                task.path.parent.mkdir(parents=True, exist_ok=True)
                LOGGER.info("Generating %s", task.path)
                task.action()
                generated.append(task.path)

            progress.update(1)

    return GenerationResult(generated=generated, skipped=skipped)


def collect_artifact_files(artifact_dir: Path) -> list[Path]:
    return sorted(path for path in artifact_dir.rglob("*") if path.is_file())


def unique_preserving_order(values: list[str]) -> list[str]:
    seen: set[str] = set()
    unique_values: list[str] = []

    for value in values:
        if value in seen:
            continue
        seen.add(value)
        unique_values.append(value)

    return unique_values


def build_upload_targets(s3_buckets: list[str], gcs_buckets: list[str]) -> list[UploadTarget]:
    targets = [UploadTarget(provider="s3", bucket=bucket) for bucket in unique_preserving_order(s3_buckets)]
    targets.extend(UploadTarget(provider="gcs", bucket=bucket) for bucket in unique_preserving_order(gcs_buckets))
    return targets


def object_name_for(prefix: str, relative_path: Path) -> str:
    normalized_prefix = prefix.strip("/")
    relative_name = relative_path.as_posix()
    return f"{normalized_prefix}/{relative_name}" if normalized_prefix else relative_name


def get_s3_client():
    client = getattr(UPLOAD_CLIENTS, "s3_client", None)
    if client is None:
        client = boto3.client("s3")
        UPLOAD_CLIENTS.s3_client = client
    return client


def get_gcs_client():
    client = getattr(UPLOAD_CLIENTS, "gcs_client", None)
    if client is None:
        client = storage.Client()
        UPLOAD_CLIENTS.gcs_client = client
    return client


def upload_single_artifact(target: UploadTarget, artifact_dir: Path, path: Path, prefix: str) -> str:
    relative_path = path.relative_to(artifact_dir)
    object_name = object_name_for(prefix, relative_path)

    if target.provider == "s3":
        get_s3_client().upload_file(str(path), target.bucket, object_name)
        return f"s3://{target.bucket}/{object_name}"

    bucket = get_gcs_client().bucket(target.bucket)
    blob = bucket.blob(object_name)
    blob.upload_from_filename(str(path))
    return f"gs://{target.bucket}/{object_name}"


def upload_artifacts(artifact_dir: Path, paths: list[Path], targets: list[UploadTarget], prefix: str, workers: int) -> None:
    if not paths:
        LOGGER.warning("No files found under %s. Nothing to upload.", artifact_dir)
        return

    for target in targets:
        LOGGER.info("Queueing %s files for upload to %s", len(paths), target.display_name)

    errors: list[str] = []
    bars = {
        target: make_progress(total=len(paths), desc=f"Upload {target.display_name}", unit="file", position=index)
        for index, target in enumerate(targets)
    }

    try:
        with ThreadPoolExecutor(max_workers=workers) as executor:
            future_to_job = {
                executor.submit(upload_single_artifact, target, artifact_dir, path, prefix): (target, path)
                for target in targets
                for path in paths
            }

            for future in as_completed(future_to_job):
                target, path = future_to_job[future]
                bars[target].update(1)

                try:
                    future.result()
                except Exception as exc:  # pragma: no cover - depends on live cloud clients.
                    object_name = object_name_for(prefix, path.relative_to(artifact_dir))
                    errors.append(f"{target.display_name}/{object_name}: {exc}")
    finally:
        for bar in bars.values():
            bar.close()

    if errors:
        error_lines = "\n".join(f"- {message}" for message in errors[:10])
        remaining = len(errors) - min(len(errors), 10)
        if remaining > 0:
            error_lines = f"{error_lines}\n- ... and {remaining} more failures"
        raise RuntimeError(f"Upload completed with {len(errors)} failure(s):\n{error_lines}")

    LOGGER.info("Uploaded %s files to %s destination(s).", len(paths), len(targets))


def resolve_modes(generate: bool, upload: bool) -> tuple[bool, bool]:
    if generate or upload:
        return generate, upload
    return True, True


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", default=str(DEFAULT_ARTIFACT_DIR), help="Artifact directory. Defaults to testdata.")
    parser.add_argument("--generate", action="store_true", help="Generate missing fixtures in the artifact directory.")
    parser.add_argument("--upload", action="store_true", help="Upload existing files from the artifact directory.")
    parser.add_argument(
        "--s3-bucket",
        action="append",
        default=[],
        help="S3 bucket destination. Repeat the flag to upload to multiple S3 buckets.",
    )
    parser.add_argument(
        "--gcs-bucket",
        action="append",
        default=[],
        help="Google Cloud Storage destination. Repeat the flag to upload to multiple GCS buckets.",
    )
    parser.add_argument("--prefix", default="", help="Optional upload prefix applied to every destination.")
    parser.add_argument("--upload-workers", type=int, default=8, help="Parallel upload workers. Defaults to 8.")
    args = parser.parse_args()

    args.generate, args.upload = resolve_modes(args.generate, args.upload)

    if args.upload_workers < 1:
        parser.error("--upload-workers must be at least 1.")

    artifact_dir = Path(args.artifact_dir).expanduser()
    if args.upload and not artifact_dir.exists() and not args.generate:
        parser.error(f"artifact directory does not exist: {artifact_dir}")

    args.artifact_dir = artifact_dir
    return args


def main() -> None:
    setup_logging()
    args = parse_args()

    generation_result = GenerationResult(generated=[], skipped=[])
    if args.generate:
        generation_result = generate_missing_fixtures(args.artifact_dir)
        LOGGER.info(
            "Generation complete for %s: %s created, %s skipped.",
            args.artifact_dir,
            len(generation_result.generated),
            len(generation_result.skipped),
        )

    if args.upload:
        targets = build_upload_targets(args.s3_bucket, args.gcs_bucket)
        if not targets:
            LOGGER.info("Upload requested but no buckets were provided. Use --s3-bucket and/or --gcs-bucket.")
        else:
            artifact_files = collect_artifact_files(args.artifact_dir)
            LOGGER.info("Found %s files under %s for upload.", len(artifact_files), args.artifact_dir)
            upload_artifacts(args.artifact_dir, artifact_files, targets, args.prefix, args.upload_workers)

    LOGGER.info(
        "Done. Artifact directory: %s. Generated: %s. Skipped existing: %s.",
        args.artifact_dir,
        len(generation_result.generated),
        len(generation_result.skipped),
    )


if __name__ == "__main__":
    main()
